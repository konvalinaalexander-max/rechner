"""
Prüfläufe für den Nachweis-Rechner.

Der Rechner ist nur so viel wert wie sein Nachweis, dass er wie Excel
rechnet. Hier stehen die Läufe, die das belegen oder widerlegen.
"""
import re, time, collections
import openpyxl
from pycel import ExcelCompiler
from xlsmpatch import Mappe

BILANZBLAETTER = ["Fläche", "D_Fläche", "Tierb", "D_Tierb", "SB1", "SB2",
                  "Gem1", "Gem2", "OB", "Allg", "FormE", "D_Gemüse1",
                  "D_Gemüse2", "D_ObstBeeren", "D_Hd", "Hd", "D_Klick", "Q"]

FUNKTION = re.compile(r"([A-ZÄÖÜ][A-ZÄÖÜ0-9\.]*)\s*\(")


def frische_mappe(quelle, ziel, blaetter):
    """Kopie ohne zwischengespeicherte Formelergebnisse."""
    m = Mappe(quelle)
    n = m.formelwerte_verwerfen(blaetter)
    m.speichern(ziel)
    return n


def formelzellen(quelle, blaetter):
    """[(Blatt, Zelle, Formel, gespeicherter Wert)]"""
    wf = openpyxl.load_workbook(quelle, data_only=False)
    wv = openpyxl.load_workbook(quelle, data_only=True)
    raus = []
    for b in blaetter:
        if b not in wf.sheetnames:
            continue
        sf, sv = wf[b], wv[b]
        for zeile in sf.iter_rows():
            for c in zeile:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    raus.append((b, c.coordinate, v, sv[c.coordinate].value))
                elif type(v).__name__ == "ArrayFormula":
                    raus.append((b, c.coordinate, str(v.text), sv[c.coordinate].value))
    return raus


def formeltreue(quelle, frisch, blaetter, grenze=1e-6, laut=True):
    """Jede Formelzelle neu rechnen und mit dem gespeicherten Wert vergleichen."""
    zellen = formelzellen(quelle, blaetter)
    exc = ExcelCompiler(frisch)
    treffer = abweichung = unrechenbar = leer = 0
    abweichungen, fehler = [], []
    t0 = time.time()
    for i, (b, co, formel, soll) in enumerate(zellen):
        try:
            ist = exc.evaluate(f"{b}!{co}")
        except Exception as e:
            unrechenbar += 1
            fehler.append((b, co, formel, type(e).__name__, str(e)[:120]))
            continue
        if soll is None and ist in (None, "", 0):
            leer += 1
            continue
        gleich = False
        if isinstance(soll, (int, float)) and isinstance(ist, (int, float)) \
                and not isinstance(soll, bool) and not isinstance(ist, bool):
            gleich = abs(float(soll) - float(ist)) <= max(grenze, abs(float(soll)) * 1e-9)
        else:
            gleich = (str(soll).strip() == str(ist).strip())
        if gleich:
            treffer += 1
        else:
            abweichung += 1
            abweichungen.append((b, co, formel, soll, ist))
        if laut and i and i % 2000 == 0:
            print(f"    ... {i}/{len(zellen)}  ({time.time()-t0:.0f}s)")
    return dict(gesamt=len(zellen), treffer=treffer, abweichung=abweichung,
                unrechenbar=unrechenbar, leer=leer,
                abweichungen=abweichungen, fehler=fehler,
                dauer=time.time() - t0)


def funktionsliste(quelle, blaetter=None):
    """Welche Excel-Funktionen kommen vor, und wie oft."""
    wf = openpyxl.load_workbook(quelle, data_only=False)
    zaehler = collections.Counter()
    blaetter = blaetter or wf.sheetnames
    for b in blaetter:
        if b not in wf.sheetnames:
            continue
        for zeile in wf[b].iter_rows():
            for c in zeile:
                v = c.value
                if type(v).__name__ == "ArrayFormula":
                    v = v.text
                if isinstance(v, str) and v.startswith("="):
                    for f in FUNKTION.findall(v):
                        zaehler[f] += 1
    return zaehler


def pycel_kennt():
    """Funktionsnamen, die pycel auswerten kann."""
    from pycel import excelutil          # noqa
    from pycel.lib import function_helpers   # noqa
    import pycel.excelformula as ef
    namen = set()
    for quelle in ("pycel.lib.function_helpers",):
        pass
    from pycel.excelformula import ExcelFormula
    try:
        namen |= set(ExcelFormula.default_modules)
    except Exception:
        pass
    import importlib
    for modul in ["pycel.lib.binary", "pycel.lib.date_time", "pycel.lib.engineering",
                  "pycel.lib.information", "pycel.lib.logical", "pycel.lib.lookup",
                  "pycel.lib.stats", "pycel.lib.text", "pycel.excellib"]:
        try:
            m = importlib.import_module(modul)
        except ImportError:
            continue
        for name in dir(m):
            obj = getattr(m, name)
            if callable(obj) and not name.startswith("_"):
                namen.add(name.upper().rstrip("_").replace("_", "."))
    return namen
