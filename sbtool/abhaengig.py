"""
Abhängigkeiten der Gesamtbilanz.

Beantwortet zwei Fragen, die man sonst nur raten kann:
  * An welchen Zellen hängt die Bilanz überhaupt?
  * Welche davon sind Eingaben - also Zellen ohne Formel, die jemand füllt?
"""
import networkx as nx
import openpyxl
from pycel import ExcelCompiler


def kegel(exc, ziele):
    """Alle Zellen, von denen die Zielzellen abhängen (einschliesslich)."""
    for z in ziele:
        exc.evaluate(z)
    g = exc.dep_graph
    knoten = {}
    for z in ziele:
        k = exc.cell_map[z]
        knoten[z] = k
        for v in nx.ancestors(g, k):
            knoten[str(v.address)] = v
    return knoten


def einteilen(exc, knoten, quelle):
    """Knoten trennen in Formelzellen, Eingabezellen und Bereiche."""
    wf = openpyxl.load_workbook(quelle, data_only=False)
    wv = openpyxl.load_workbook(quelle, data_only=True)
    formeln, eingaben, bereiche = [], [], []
    for adr, k in knoten.items():
        if ":" in adr:
            bereiche.append(adr)
            continue
        if "!" not in adr:
            continue
        blatt, co = adr.split("!", 1)
        if blatt not in wf.sheetnames:
            continue
        c = wf[blatt][co]
        v = c.value
        istformel = (isinstance(v, str) and v.startswith("=")) or type(v).__name__ == "ArrayFormula"
        if istformel:
            formeln.append(adr)
        else:
            eingaben.append((adr, wv[blatt][co].value, c.protection.locked))
    return formeln, eingaben, bereiche


def bereiche_aufloesen(exc, knoten, quelle):
    """Zellen, die nur über einen Bereich in den Kegel kommen, einzeln nachtragen."""
    from pycel.excelutil import AddressRange
    wf = openpyxl.load_workbook(quelle, data_only=False)
    wv = openpyxl.load_workbook(quelle, data_only=True)
    extra_formeln, extra_eingaben = [], []
    for adr in list(knoten):
        if ":" not in adr:
            continue
        try:
            ar = AddressRange(adr)
        except Exception:
            continue
        blatt = ar.sheet
        if blatt not in wf.sheetnames:
            continue
        for zeile in ar.rows:
            for a in zeile:
                co = a.coordinate
                voll = f"{blatt}!{co}"
                if voll in knoten:
                    continue
                c = wf[blatt][co]
                v = c.value
                if v is None:
                    continue
                istformel = (isinstance(v, str) and v.startswith("=")) or \
                            type(v).__name__ == "ArrayFormula"
                (extra_formeln if istformel else extra_eingaben).append(
                    voll if istformel else (voll, wv[blatt][co].value, c.protection.locked))
    return extra_formeln, extra_eingaben
