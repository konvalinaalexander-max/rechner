"""
Warn- und Prüfzellen der Mappe.

Die Mappe kontrolliert sich selbst: einzelne Zellen zeigen einen Hinweistext,
wenn etwas nicht stimmt ("LN kontrollieren!", "Anbaufläche stimmt nicht mit
dem Blatt Fläche überein!"). Diese Zellen sind über die ganze Mappe verstreut.
Hier werden sie eingesammelt und ihr Zustand gemeldet.
"""
import openpyxl

# Texte, die als Warnung gelten (Ausschnitte genügen)
WARNWORTE = ["kontrollieren", "stimmt nicht", "zu viele", "zuviele", "nicht zulässig",
             "unzulässig", "fehlt", "Fehler", "prüfen", "pruefen", "überschritten",
             "ausserhalb", "nicht möglich", "Achtung", "Warnung", "!"]


def texte_der_mappe(pfad):
    """Alle Hinweistexte, die die Mappe kennt (Blatt Texte)."""
    wv = openpyxl.load_workbook(pfad, data_only=True)
    raus = {}
    if "Texte" in wv.sheetnames:
        ws = wv["Texte"]
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and len(v) > 8:
                raus[f"Texte!A{r}"] = v
    return raus


def kandidaten(pfad):
    """Formelzellen, deren Formel auf einen Warntext verweist."""
    wf = openpyxl.load_workbook(pfad, data_only=False)
    texte = texte_der_mappe(pfad)
    warn_adr = {a for a, t in texte.items()
                if any(w.lower() in t.lower() for w in WARNWORTE)}
    warn_ref = {a.split("!")[1] for a in warn_adr}
    raus = []
    for b in wf.sheetnames:
        if b.startswith("Dialog"):
            continue
        for zeile in wf[b].iter_rows():
            for c in zeile:
                v = c.value
                if type(v).__name__ == "ArrayFormula":
                    v = v.text
                if not isinstance(v, str) or not v.startswith("="):
                    continue
                if "Texte!" in v:
                    for ref in warn_ref:
                        if f"Texte!{ref}" in v or f"Texte!${ref[0]}${ref[1:]}" in v:
                            raus.append((b, c.coordinate, v))
                            break
    return raus, texte


def zustand(pfad, kand):
    """Welche Warnzellen zeigen gerade einen Text?"""
    wv = openpyxl.load_workbook(pfad, data_only=True)
    aktiv, ruhig, unbekannt = [], [], []
    for b, co, formel in kand:
        v = wv[b][co].value
        if isinstance(v, str) and v.strip() and v.strip() not in ("0", " "):
            aktiv.append((b, co, v, formel))
        elif v is None:
            unbekannt.append((b, co, formel))
        else:
            ruhig.append((b, co, v))
    return aktiv, ruhig, unbekannt


# --- direkter Weg: Zellen, deren angezeigter Wert nach einer Warnung aussieht ---
MARKER = ["kontrollieren", "stimmt nicht", "zu viele", "zuviele", "überein!",
          "nicht zulässig", "unzulässig", "überschritten", "ausserhalb",
          "nicht möglich", "zu hoch", "zu tief", "Fehler:", "prüfen!"]


def sichtbare_warnungen(pfad, blaetter=None):
    """Alle Zellen der Mappe, deren gespeicherter Wert ein Warntext ist."""
    wv = openpyxl.load_workbook(pfad, data_only=True)
    wf = openpyxl.load_workbook(pfad, data_only=False)
    raus = []
    for b in (blaetter or wv.sheetnames):
        if b.startswith("Dialog") or b.startswith("D_") or b in ("Texte", "Q"):
            continue
        ws, wsf = wv[b], wf[b]
        for zeile in ws.iter_rows():
            for c in zeile:
                v = c.value
                if not isinstance(v, str) or len(v) < 6:
                    continue
                if any(m.lower() in v.lower() for m in MARKER):
                    f = wsf[c.coordinate].value
                    if isinstance(f, str) and f.startswith("="):
                        raus.append((b, c.coordinate, v, f))
    return raus
